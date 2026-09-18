# -*- coding: utf-8 -*-
"""生成《MusicPlayer 操作手册》Word 文档（.docx）。
云盘连接章节为重点，所有需要截图的位置以带边框灰底占位框标注。
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = r"D:\我的文件\Documents\WorkBuddy\音乐播放器"
OUT = os.path.join(ROOT, "MusicPlayer-操作手册-v0.1.0.docx")
CJK = "微软雅黑"

doc = Document()

# ---- 默认正文字体（中英文）----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.25

def set_heading_font(style_name, size, color):
    st = doc.styles[style_name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(*color)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

set_heading_font("Title", 26, (0x1F, 0x4E, 0x79))
set_heading_font("Heading 1", 16, (0x1F, 0x4E, 0x79))
set_heading_font("Heading 2", 13, (0x2E, 0x5C, 0x8A))
set_heading_font("Heading 3", 11.5, (0x3A, 0x3A, 0x3A))

def shade(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)

def add_callout(label, text, fill="FFF4CE", label_rgb=(0x8A, 0x6D, 0x00)):
    """提示 / 注意 / 警告 高亮框。"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "C9A227")
    pbdr.append(left)
    pPr.append(pbdr)
    shade(p, fill)
    r = p.add_run(label + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor(*label_rgb)
    r2 = p.add_run(text)
    return p

def add_screenshot(caption, height_cm=5.5):
    """带边框灰底占位框 + 图注；用户可点进单元格插入截图。"""
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    # 高度
    tr = cell._tc.getparent()
    trPr = tr.get_or_add_trPr()
    trH = OxmlElement("w:trHeight")
    trH.set(qn("w:val"), str(int(height_cm * 567)))
    trH.set(qn("w:hRule"), "atLeast")
    trPr.append(trH)
    # 灰底
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F2F2F2")
    tcPr.append(shd)
    # 占位文字
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("【截图占位：" + caption + "】")
    run.font.color.rgb = RGBColor(0x9A, 0x9A, 0x9A)
    run.font.size = Pt(10)
    # 图注
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run("图：" + caption)
    cr.italic = True
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    cap.paragraph_format.space_after = Pt(8)

def step(n, text):
    p = doc.add_paragraph(style="List Number")
    p.add_run(text)
    return p

def bullet(text):
    return doc.add_paragraph(text, style="List Bullet")

# ============================ 封面 ============================
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("MusicPlayer 操作手册")
r.bold = True

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run("Windows 桌面版 · 自包含便携版 v0.1.0")
sr.font.size = Pt(12)
sr.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

doc.add_paragraph()
add_callout("说明",
            "本手册基于 v0.1.0 编写。云盘连接功能当前仅支持「坚果云 WebDAV」。"
            "文档中所有标注【截图占位】的方框，请自行截取对应界面后替换为真实图片。")

doc.add_page_break()

# ============================ 目录式概览（手动） ============================
doc.add_heading("目录", level=1)
for i, s in enumerate([
    "1. 软件简介", "2. 系统要求", "3. 安装与启动",
    "4. 界面与基本操作", "5. 连接云盘（重点）", "6. 设置说明",
    "7. 歌词功能", "8. 数据与隐私", "9. 常见问题（FAQ）", "10. 合规声明",
], 1):
    doc.add_paragraph(s)

doc.add_page_break()

# ============================ 1. 软件简介 ============================
doc.add_heading("1. 软件简介", level=1)
doc.add_paragraph(
    "MusicPlayer 是一款跨平台音乐播放器，本手册针对其 Windows 桌面端（基于 WPF + .NET 10 + "
    "BASS 音频引擎）。主要功能包括：本地曲库管理、播放列表、歌词显示，以及通过「坚果云 WebDAV」"
    "连接个人云盘，直接在云端音乐上播放与同步播放列表。")
doc.add_paragraph("本版本为「自包含便携版」：解压即用，无需安装 .NET Runtime；"
                  "曲库、设置、播放列表等数据随程序目录存放，便于携带。")

# ============================ 2. 系统要求 ============================
doc.add_heading("2. 系统要求", level=1)
bullet("操作系统：Windows 10 版本 19041 或更高（仅 x64 / 64 位）。")
bullet("不需要预先安装 .NET Runtime（程序已自带运行时）。")
bullet("建议：稳定的网络（使用云盘功能时）。")
bullet("可选：坚果云账号（用于云盘播放 / 同步）。")

# ============================ 3. 安装与启动 ============================
doc.add_heading("3. 安装与启动", level=1)
step(1, "从发布页下载压缩包 MusicPlayer-Desktop-win-x64-v0.1.0.zip。")
step(2, "将其解压到任意「可写目录」（例如 D:\\MusicPlayer）。"
        "注意不要解压到 C:\\Program Files 等需要管理员权限的目录，否则数据无法写入。")
step(3, "进入解压后的文件夹，双击 MusicPlayer.Desktop.exe 启动程序。")
step(4, "首次启动若被系统拦截（见下方提示），按提示放行即可。")
add_screenshot("解压后的程序文件夹与 MusicPlayer.Desktop.exe 入口", height_cm=4.5)
add_callout("注意（SmartScreen 拦截）",
            "本程序未做代码签名。首次在他人电脑上运行时，Windows SmartScreen 可能弹出"
            "「Windows 已保护你的电脑」。点击「更多信息」→「仍要运行」即可继续。"
            "这是正常现象，不影响功能。",
            fill="FDECEA", label_rgb=(0xB0, 0x2A, 0x1E))

# ============================ 4. 界面与基本操作 ============================
doc.add_heading("4. 界面与基本操作", level=1)
doc.add_paragraph("启动后主界面大致分为：顶部菜单栏、播放控制区、播放列表区、右侧歌词区。")
add_screenshot("主界面总览（菜单栏 / 播放控制 / 播放列表 / 歌词区）", height_cm=6)
doc.add_heading("4.1 添加本地音乐", level=2)
step(1, "通过菜单「文件 / 导入」或界面上的导入按钮，选择音乐文件夹或文件加入曲库。")
step(2, "添加后，曲目出现在左侧曲库 / 播放列表中，双击即可播放。")
doc.add_heading("4.2 播放控制", level=2)
bullet("播放 / 暂停、上一首 / 下一首、进度拖动、音量调节。")
bullet("右侧歌词区随播放进度高亮当前行（需有对应歌词，详见第 7 章）。")

# ============================ 5. 连接云盘（重点） ============================
doc.add_heading("5. 连接云盘（重点）", level=1)
doc.add_paragraph(
    "MusicPlayer 支持连接你自己的坚果云网盘，直接在云端音乐上播放，并实现播放列表的跨设备同步。"
    "连接方式采用坚果云提供的 WebDAV 协议。本章给出完整步骤，所有关键界面均留截图占位。")

doc.add_heading("5.1 前置准备：获取坚果云「应用密码」", level=2)
add_callout("重要",
            "坚果云 WebDAV 不能使用你的「登录密码」，必须使用专门的「应用密码（第三方应用授权密码）」。"
            "若误填登录密码，连接会失败。",
            fill="FDECEA", label_rgb=(0xB0, 0x2A, 0x1E))
step(1, "用浏览器登录坚果云网页端（https://www.jianguoyun.com/）。")
step(2, "进入「账号信息」→「安全选项」。")
step(3, "找到「第三方应用授权 / 应用密码」区域，点击「生成应用密码」。")
step(4, "按提示验证后，复制生成的一串应用密码（仅显示一次，请妥善保存）。")
add_screenshot("坚果云网页端：安全选项 → 应用密码生成界面（红框标出「生成应用密码」按钮与生成的密码）", height_cm=6)

doc.add_heading("5.2 打开「云盘」窗口", level=2)
step(1, "在主界面顶部菜单栏，点击「云盘」菜单项。")
step(2, "弹出「云盘」窗口，顶部即为「连接云盘（坚果云 WebDAV）」表单区。")
add_screenshot("主界面菜单栏：点击「云盘」菜单项（用箭头标出菜单位置）", height_cm=4.5)
add_screenshot("「云盘」窗口：连接表单与浏览区总览", height_cm=6)

doc.add_heading("5.3 填写连接信息", level=2)
doc.add_paragraph("在「连接云盘（坚果云 WebDAV）」表单中逐项填写：")
bullet("名称：自定义显示名，例如「我的坚果云」。")
bullet("URL：WebDAV 根地址，固定为 https://dav.jianguoyun.com/dav/ （注意末尾的斜杠 / 必须保留）。")
bullet("用户名：你的坚果云账号（注册邮箱或用户名）。")
bullet("密码：5.1 中生成的「应用密码」（不是登录密码）。")
add_screenshot("连接表单填写示例（名称 / URL / 用户名 / 应用密码，URL 末尾斜杠与密码字段高亮标出）", height_cm=5)

doc.add_heading("5.4 连接并添加", level=2)
step(1, "确认信息无误后，点击「连接并添加」按钮。")
step(2, "窗口下方状态栏会显示结果：成功时提示已连接，下方「已连接来源」列表出现该来源，"
        "浏览区加载出云盘根目录。")
step(3, "连接成功后凭据会被保存在本机（见第 8 章），下次启动会自动重连，无需重复填写。")
add_screenshot("连接成功：状态栏提示「已连接」+「已连接来源」列表 + 根目录浏览内容", height_cm=5)
add_callout("连接失败排查",
            "若提示连接失败，请依次检查：① URL 末尾是否带「/」；② 密码是否为「应用密码」而非登录密码；"
            "③ 坚果云账号是否已开通 WebDAV（默认开通）；④ 本机网络是否可访问 dav.jianguoyun.com。",
            fill="FFF4CE", label_rgb=(0x8A, 0x6D, 0x00))

doc.add_heading("5.5 浏览与添加云端音乐", level=2)
doc.add_paragraph("连接成功后即可在浏览区操作：")
bullet("双击文件夹：进入下一层目录。")
bullet("双击音乐文件：将其加入播放列表（不自动播放）。")
bullet("右键曲目：可「播放」「加入播放列表」「进入文件夹」「加入此文件夹音乐」。")
bullet("选中条目后，点「加入选中到列表」批量加入。")
add_screenshot("云盘浏览区：文件夹进入 + 选中音乐加入播放列表", height_cm=5.5)

doc.add_heading("5.6 播放列表同步（保存到云盘 / 从云盘加载）", level=2)
doc.add_paragraph(
    "播放列表可与云盘互相同步，便于在多台设备间共享同一份歌单（列表中记录的是云盘相对路径，"
    "在另一台已连同一云盘的设备上可还原）。")
bullet("「保存到云盘」：将当前播放列表上传到云盘。")
bullet("「从云盘加载」：从云盘下载此前保存的播放列表。")
bullet("这两个按钮位于主界面播放列表区域，以及「云盘」窗口浏览区工具栏。")
add_screenshot("主界面播放列表区：「保存到云盘」「从云盘加载」按钮位置", height_cm=4.5)

doc.add_heading("5.7 常见问题（云盘）", level=2)
bullet("状态栏提示「请先连接一个云盘来源」：说明还没有已连接的云盘，请先完成 5.2–5.4。")
bullet("「云盘连接已失效，请重新连接」：凭据过期或网络变化，重新打开「云盘」窗口并连接即可。")
bullet("能连上但无法播放：确认该文件为支持的音频格式，且网络可访问。")

# ============================ 6. 设置说明 ============================
doc.add_heading("6. 设置说明", level=1)
doc.add_paragraph("主界面菜单「设置」可打开设置窗口，常用项如下：")
bullet("下载 / 缓存目录：云端音乐会先缓存到此处再播放，可点「浏览」修改；右侧显示已用空间。")
bullet("歌词：可指定额外的本地歌词目录；「允许网络获取歌词」默认关闭（隐私优先）；"
       "「把歌词上传回云盘同目录」可让歌词跟着歌曲走。")
bullet("ReplayGain（响度归一）：统一不同曲目的音量，避免忽大忽小。")
bullet("音量：总音量调节滑块。")
add_screenshot("设置窗口：缓存目录 / 歌词选项 / 响度归一 / 音量", height_cm=6)

# ============================ 7. 歌词功能 ============================
doc.add_heading("7. 歌词功能", level=1)
doc.add_paragraph("歌词来源按以下优先级匹配（可在设置中调整）：")
step(1, "本地 .lrc 文件 / 音频内嵌歌词。")
step(2, "云盘同目录下的 .lrc 文件（已连接云盘时）。")
step(3, "网络歌词 API（默认关闭；在设置中开启后才联网获取）。")
doc.add_paragraph("纯文本歌词（无时间轴）不会高亮；带时间轴的 .lrc 会随播放进度逐行高亮。"
                  "若开启「把歌词上传回云盘同目录」，拿到歌词后会写回云盘，使歌词跟随歌曲。")

# ============================ 8. 数据与隐私 ============================
doc.add_heading("8. 数据与隐私", level=1)
bullet("数据位置：曲库、设置、播放列表保存在程序目录下的 library.db（真便携，换电脑随身走）。")
bullet("云盘凭据：坚果云账号与「应用密码」保存在本机 library.db 中。"
      "当前为明文存储，请注意使用环境安全；后续版本计划加入系统级加密（DPAPI）。")
bullet("下载缓存：默认位于「音乐库\\MusicPlayerCache」，可在设置中更改。")
add_callout("隐私提醒",
            "云盘「应用密码」等同于你云盘的部分访问权限，请勿在公共或不信任的电脑上连接个人云盘；"
            "离开前建议在坚果云后台吊销该应用密码。",
            fill="FFF4CE", label_rgb=(0x8A, 0x6D, 0x00))

# ============================ 9. FAQ ============================
doc.add_heading("9. 常见问题（FAQ）", level=1)
bullet("没声音：检查系统默认输出设备；确认解压目录包含 BASS 原生文件（bass.dll 等）；"
       "必要时查看程序目录下的 diag.log 定位问题。")
bullet("首次运行被杀软拦截：见 3.4 的 SmartScreen 处理。")
bullet("云盘连不上：见 5.4 / 5.7 的排查清单。")
bullet("数据在哪：见第 8 章；换电脑时连同整个程序文件夹一起拷贝即可。")

# ============================ 10. 合规声明 ============================
doc.add_heading("10. 合规声明", level=1)
bullet("本软件仅用于播放用户自己拥有合法权利的音乐，仅支持连接用户「自有」的坚果云网盘。")
bullet("不提供、也不鼓励任何第三方资源搜索 / 分享 / 下载功能。")
bullet("音频引擎使用 BASS（Un4seen），非商业用途免费；若公开发布请遵守其授权条款。")

doc.add_paragraph()
foot = doc.add_paragraph()
foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = foot.add_run("— MusicPlayer v0.1.0 操作手册 · 截图待补充 —")
fr.font.size = Pt(9)
fr.font.color.rgb = RGBColor(0x99, 0x99, 0x99)

doc.save(OUT)
print("已生成：", OUT)
print("段落数：", len(doc.paragraphs))
